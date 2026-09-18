"""Сборка презентации решения из шаблона ЛЦТ2026 (python-pptx).

Слайды шаблона переставляются в нужном порядке, лишние удаляются, тексты вписываются
с сохранением стилей шаблона, диаграммы получают наши данные.
Запуск: .venv/Scripts/python build_deck.py → Geoscan_fleet_planner.pptx
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"
OUT = HERE / "Geoscan_fleet_planner.pptx"

TEAM = "[НАЗВАНИЕ КОМАНДЫ]"

# порядок: номера слайдов шаблона (1-based)
ORDER = [7, 8, 9, 10, 11, 24, 25, 14, 16, 21, 27, 19, 22, 20, 18, 15, 17, 12]


# ---------------------------------------------------------------- утилиты
def shapes_by_ph(slide) -> dict[int, object]:
    return {sh.placeholder_format.idx: sh for sh in slide.placeholders}


_PPR_TAIL = ("tabLst", "defRPr", "extLst")


def no_bullet(p):
    """Абзац без маркера и отступа (для подзаголовков и коротких подписей)."""
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", "0")
    ppr.set("indent", "0")
    for tag in ("buNone", "buAutoNum", "buChar", "buBlip"):
        for el in ppr.findall(qn(f"a:{tag}")):
            ppr.remove(el)
    bu = ppr.makeelement(qn("a:buNone"), {})
    tail = next((c for c in ppr if c.tag.split("}")[1] in _PPR_TAIL), None)
    if tail is not None:
        tail.addprevious(bu)
    else:
        ppr.append(bu)


def fill(shape, paras, size=None, bold_first=False, color=None, bullets=True):
    """Вписать абзацы, сохранив оформление первого абзаца/ранa шаблона.
    paras: список строк или (строка, {'bold':..,'size':..,'bullet':..}).
    Жирные строки (подзаголовки) — всегда без маркера."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr = None
    if p0.runs and p0.runs[0]._r.rPr is not None:
        rpr = copy.deepcopy(p0.runs[0]._r.rPr)
    # В пустых плейсхолдерах шаблона цвет «при наборе» хранится в endParaRPr (по умолчанию
    # у макета белый текст) — переносим его в оформление вписываемого текста.
    end = p0._p.find(qn("a:endParaRPr"))
    if end is not None:
        if rpr is None:
            rpr = copy.deepcopy(end)
            rpr.tag = qn("a:rPr")
        elif rpr.find(qn("a:solidFill")) is None and end.find(qn("a:solidFill")) is not None:
            rpr.insert(0, copy.deepcopy(end.find(qn("a:solidFill"))))
    # удалить все абзацы кроме первого, очистить первый
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0._p):
        if r.tag.endswith("}r") or r.tag.endswith("}br") or r.tag.endswith("}fld"):
            p0._p.remove(r)
    for i, item in enumerate(paras):
        text, opts = (item, {}) if isinstance(item, str) else item
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and ppr is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, copy.deepcopy(ppr))
        run = p.add_run()
        if rpr is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, copy.deepcopy(rpr))
        run.text = text
        b = opts.get("bold", bold_first and i == 0)
        if b:
            run.font.bold = True
        s = opts.get("size", size)
        if s:
            run.font.size = Pt(s)
        if color or opts.get("color"):
            from pptx.dml.color import RGBColor

            run.font.color.rgb = RGBColor.from_string(opts.get("color", color))
        if not opts.get("bullet", bullets) or b:
            no_bullet(p)
        if opts.get("space_before"):
            p.space_before = Pt(opts["space_before"])


def title(slide, text, pill_pad=0.62, per_char=0.215):
    """Заголовок + подгонка ширины «плашки» под длину заголовка."""
    ph = shapes_by_ph(slide).get(0)
    if ph is None:
        return
    fill(ph, [text])
    for sh in slide.shapes:
        if (sh.shape_type == 1 and "Скругленный" in sh.name and abs(Emu(sh.top).inches - 0.35) < 0.05
                and abs(Emu(sh.height).inches - 0.68) < 0.05):
            sh.width = Inches(min(12.6, pill_pad + per_char * len(text)))


def text_boxes(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder]


def find_tb(slide, startswith):
    for sh in text_boxes(slide):
        if sh.text_frame.text.strip().startswith(startswith):
            return sh
    raise KeyError(startswith)


def picture(slide, idx, path):
    """Картинка в плейсхолдер (с обрезкой по его пропорциям) на месте плейсхолдера слайда."""
    ph = shapes_by_ph(slide)[idx]
    box = (ph.left, ph.top, ph.width, ph.height)
    pic = ph.insert_picture(str(path))
    pic.left, pic.top, pic.width, pic.height = box
    set_alt(pic, path)
    return pic


def picture_fit(slide, idx, path, max_w=None):
    """Картинка без обрезки: вписать в рамку плейсхолдера по центру, плейсхолдер удалить."""
    from PIL import Image

    ph = shapes_by_ph(slide)[idx]
    left, top, w, h = ph.left, ph.top, ph.width, ph.height
    ph._element.getparent().remove(ph._element)
    iw, ih = Image.open(path).size
    if max_w:
        w = min(w, Inches(max_w))
    scale = min(w / iw, h / ih)
    pw, phh = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(path), left + (ph.width - pw) // 2 if not max_w else left,
                                   top + (h - phh) // 2, pw, phh)
    set_alt(pic, path)
    return pic


ALT = {
    "demo_wide.jpg": "Интерфейс сервиса: фронт Парето, график вылетов и маршруты четырёх бортов на карте",
    "demo_portrait.jpg": "Маршруты четырёх бортов вокруг запретной зоны в демо-сценарии",
    "ui_demo.jpg": "Интерфейс: фронт Парето, показатели плана и маршруты на карте",
    "ui_large.jpg": "Интерфейс: план для 110 км² с пятью бортами и двумя ВПП",
    "geo.jpg": "Магнитная съёмка двумя Геоскан 401 с секущими маршрутами",
    "lidar.jpg": "LiDAR-съёмка двумя Геоскан 401",
    "large.jpg": "Разбиение 110 км² между самолётами и мультироторами с учётом дальности",
    "s6_14_3.1_6.5.png": "Логотип GEOSCAN",
}


def set_alt(pic, path):
    pic._element.nvPicPr.cNvPr.set("descr", ALT.get(Path(path).name, ""))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_slide_number(slide, n):
    for sh in slide.shapes:
        if sh.has_text_frame and "Номер слайда" in sh.name:
            if sh.text_frame.text.strip().isdigit():
                fill(sh, [str(n)])


# ---------------------------------------------------------------- сборка
prs = Presentation(HERE / "template.pptx")
sld_ids = prs.slides._sldIdLst
all_ids = list(sld_ids)
keep = [all_ids[i - 1] for i in ORDER]
for el in all_ids:
    sld_ids.remove(el)
    if el not in keep:
        prs.part.drop_rel(el.rId)
for el in keep:
    sld_ids.append(el)
S = list(prs.slides)
assert len(S) == len(ORDER)

# 1. Титульный
s = S[0]
ph = shapes_by_ph(s)
fill(ph[0], [f"КОМАНДА «{TEAM.strip('[]')}»" if not TEAM.startswith("[") else TEAM])
fill(ph[12], ["Задача 5. Сервис планирования и распределения беспилотных авиационных работ"])
picture_fit(s, 11, HERE / "logos" / "s6_14_3.1_6.5.png", max_w=3.4)
notes(s, "Представить команду и задачу: сервис, который распределяет авиационные работы между бортами парка Геоскан и строит каждому полётное задание.")

# 2. О команде и решении
s = S[1]
fill(shapes_by_ph(s)[0], [TEAM])
fill(find_tb(s, "В чем суть"), [
    "Веб-сервис делит область съёмки между бортами парка Геоскан и строит каждому полётное "
    "задание: галсы, развороты, вылеты по заряду, обход запретных зон, ветер и рельеф. "
    "Экспорт — KML и GeoJSON.",
])
fill(find_tb(s, "Что делает"), [
    "Geoscan Planner, Mission Planner и QGroundControl планируют один борт. Мы распределяем "
    "работу по разнородному парку и даём выбрать компромисс «время работ — налёт».",
])
fill(find_tb(s, "Капитан"), [
    "Капитан: [ФИО, специальность]",
    "Участников: [N] человек",
    "Как образовалась команда: [заполнить]",
    "Место работы / учёбы: [заполнить]",
    "Город и регион: [заполнить]",
])
picture(s, 10, IMG / "demo_wide.jpg")
notes(s, "Суть и уникальность в двух фразах. Поля в квадратных скобках заполнить данными команды.")

# 3. Состав команды
s = S[2]
title(s, "СОСТАВ КОМАНДЫ")
for sh in [find_tb(s, "Имя Фамилия")] + [x for x in text_boxes(s) if x.text_frame.text.strip() == "Имя Фамилия"]:
    fill(sh, ["[Имя Фамилия]"])
for sh in [x for x in s.shapes if x.has_text_frame and x.text_frame.text.startswith("Роль в команде")]:
    fill(sh, ["[Роль]", "[Telegram]", "[Телефон]", "[Место работы/учёбы]"])
notes(s, "Заполнить карточки участников; лишние карточки удалить вместе с рамкой и фото.")

# 4. Работа над задачей
s = S[3]
title(s, "КАК МЫ РАБОТАЛИ")
fill(shapes_by_ph(s)[27], ["[Как собрались, участвовали ли вместе в хакатонах — заполнить]"])
fill(find_tb(s, "Что вас вдохновило"), [
    "Один борт давно умеют планировать Geoscan Planner и Mission Planner. Распределять работу "
    "по разнородному парку с учётом заряда и ветра — нет. Это настоящая задача оптимизации.",
])
fill(find_tb(s, "Расскажите о самых"), [
    "Самолёту негде развернуться у запретной зоны — оставили запас на петлю и отдали эти полосы "
    "мультироторам. На 110 км² коптеры не долетали до своих полос — добавили разбиение с учётом "
    "дальности. Бенчмарк показал проигрыш равному делению — добавили мультистарт.",
])
notes(s, "Первый блок — история команды, заполнить. Сложности — реальные, из разработки.")

# 5. Коротко о решении
s = S[4]
ph = shapes_by_ph(s)
title(s, "КОРОТКО О РЕШЕНИИ")
fill(ph[49], ["Техническая суть"])
fill(ph[38], [
    "Ядро на Python рассчитывает параметры съёмки по ТТХ бортов, строит галсы с учётом ветра, "
    "развороты по Дубинсу, обходит запретные зоны, режет маршрут на вылеты по заряду и делит "
    "область между бортами.",
    "Два критерия и фронт Парето. Рельеф Copernicus DEM. Веб-интерфейс на React + Leaflet, "
    "один Docker-контейнер, ключи не нужны.",
])
fill(ph[51], ["Для кого и зачем"])
fill(ph[42], [
    "Операторы парков БВС в геодезии, мониторинге инфраструктуры, агро и геофизике: вместо "
    "ручной нарезки области на несколько бортов — готовые задания за секунды.",
    "Быстрее получить данные: на смешанном парке план до 2,9 раза быстрее равного деления "
    "области без учёта ТТХ.",
])
notes(s, "Техническая и прикладная суть. Цифра 2,9× — из бенчмарка (сценарий strong_wind).")

# 6. Проблема и решение
s = S[5]
ph = shapes_by_ph(s)
title(s, "ПРОБЛЕМА И РЕШЕНИЕ")
fill(ph[26], [("Проблема", {"bold": True}),
              "Парк из разных бортов: самолёт на 3 часа и коптеры на 40 минут. Делить область "
              "между ними вручную долго, результат далёк от оптимума, легко забыть про ветер "
              "и заряд."])
fill(ph[31], [("Решение", {"bold": True}),
              "Сервис сам делит область по производительности и дальности бортов, строит галсы, "
              "развороты и вылеты, обходит запретные зоны и проверяет аварийный уход на площадки."])
fill(ph[32], [("Результат", {"bold": True}),
              "Индивидуальные задания в KML и GeoJSON, выбор критерия «время работ» или «налёт» "
              "и набор компромиссов между ними. Расчёт — секунды."])

# 7. Как работает
s = S[6]
ph = shapes_by_ph(s)
title(s, "КАК РАБОТАЕТ СЕРВИС")
steps = [
    ("Исходные данные", "Область, разрешённая зона, NFZ, ВПП, резервные площадки, парк, тип съёмки, ветер"),
    ("Параметры съёмки", "Высота, полоса, шаг галсов и скорость из GSD, плотности LiDAR или шага геофизики"),
    ("Распределение", "Доли области по производительности и дальности бортов, балансировка по времени"),
    ("Маршруты и вылеты", "Галсы по ветру, развороты Дубинса, обход NFZ, нарезка по заряду"),
    ("Проверка и экспорт", "Уход на площадки, радиоканал, покрытие; KML и GeoJSON на каждый борт"),
]
for (h, d), (ih, idd) in zip(steps, [(26, 27), (28, 29), (30, 31), (32, 33), (34, 35)]):
    fill(ph[ih], [h], bullets=False)
    fill(ph[idd], [d], bullets=False)

# 8. Демо
s = S[7]
ph = shapes_by_ph(s)
title(s, "ДЕМО: 4 БОРТА, 6,3 КМ²")
fill(ph[14], bullets=False, paras=[
    ("Сценарий", {"bold": True}),
    "Геоскан 201, два Gemini и 801, две ВПП, запретная зона в центре, ветер 5 м/с, RGB 3 см/пикс.",
    ("Что сделал сервис", {"bold": True, "space_before": 8}),
    "Самолёт снимает основную часть длинными галсами с петлями на разворотах. Кольцо вокруг "
    "запретной зоны, где самолёту не развернуться, забрал Геоскан 801.",
    ("Итог", {"bold": True, "space_before": 8}),
    "Все работы за 58 мин, покрытие 100 %, ни одного пересечения запретной зоны. Один лучший "
    "борт потратил бы 83 мин и оставил бы неснятыми 4 % у запретной зоны.",
])
picture(s, 10, IMG / "demo_portrait.jpg")
notes(s, "Показать живое демо: сценарий demo_basic → «Рассчитать план» → «Сравнить варианты».")

# 9. Алгоритм
s = S[8]
ph = shapes_by_ph(s)
title(s, "АЛГОРИТМ")
cols = [
    ("Галсы и ветер", "Направление галсов перебираем по рёбрам полигона и оцениваем по времени с учётом "
                      "ветра. Дыры и вогнутости — клеточная декомпозиция boustrophedon."),
    ("Развороты", "Самолёт — кратчайшие пути Дубинса, R = V²/(g·tg φ): при узком шаге получается "
                  "петля, как в Geoscan Planner. Коптер — остановка и разгон."),
    ("Обход зон и заряд", "Перелёты — по графу видимости вокруг NFZ. Маршрут режется на вылеты "
                          "по заряду с резервом, длинный галс делится."),
    ("Парк и критерии", "Доли области по производительности и дальности, балансировка по факту, "
                        "локальный поиск. Фронт Парето: взвешенная сумма и ε-ограничения."),
]
for i, (h, d) in enumerate(cols):
    fill(ph[49 + i], [f"0{i + 1}"])
    fill(ph[37 + 2 * i], [h], bullets=False)
    fill(ph[38 + 2 * i], [d], bullets=False)
notes(s, "Источники: Choset 2000, Huang 2001, Coombes 2017, Dubins 1957, Lozano-Pérez 1979, "
         "Beasley 1983, Maza & Ollero 2007, DARP 2017, Mavrotas 2009 — docs/algorithm.md.")

# 10. Результаты
s = S[9]
ph = shapes_by_ph(s)
title(s, "РЕЗУЛЬТАТЫ")
chart = next(sh for sh in s.shapes if sh.has_chart).chart
cd = CategoryChartData()
cd.categories = ["6 км², 4 борта", "Геофизика", "LiDAR", "Ветер 11 м/с", "110 км²"]
cd.add_series("Один лучший борт", (1.42, 1.80, 1.74, 1.16, 2.17))
cd.add_series("Поровну без учёта ТТХ", (2.50, 1.03, 1.04, 2.92, 1.00))
chart.replace_data(cd)
chart.has_title = False
cap = s.shapes.add_textbox(Inches(0.58), Inches(1.25), Inches(6.0), Inches(0.5))
cap.text_frame.word_wrap = True
cap.text_frame.text = "Во сколько раз дольше идут работы при базовых подходах (наш план = 1)"
for r in cap.text_frame.paragraphs[0].runs:
    r.font.size = Pt(13)
    from pptx.dml.color import RGBColor
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
BIG = dict(size=24, bold_first=True, bullets=False)
fill(ph[21], ["до 2,9×"], **BIG)
fill(ph[18], ["дольше работает парк, если делить область поровну без учёта ТТХ бортов"], bullets=False)
fill(ph[22], ["1,2–2,2×"], **BIG)
fill(ph[23], ["выигрыш по времени работ против одного лучшего борта"], bullets=False)
fill(ph[24], ["98–100 %"], **BIG)
fill(ph[25], ["покрытие во всех сценариях; 0 пересечений запретных зон, 31 автотест"], bullets=False)
notes(s, "Диаграмма: во сколько раз дольше работы при базовых подходах, чем по нашему плану (=1). "
         "Данные — docs/benchmark.md, воспроизводится python bench/benchmark.py.")

# 11. Интерфейс
s = S[10]
ph = shapes_by_ph(s)
title(s, "ИНТЕРФЕЙС")
picture(s, 14, IMG / "ui_demo.jpg")
picture(s, 18, IMG / "ui_large.jpg")
fill(ph[15], bullets=False, paras=[("Фронт Парето и график вылетов", {"bold": True}),
              "Каждая точка — недоминируемый план: левее быстрее, ниже меньше налёт. Щелчок "
              "открывает план на карте."])
fill(ph[16], bullets=False, paras=[("110 км², 5 бортов, 2 ВПП", {"bold": True}),
              "Самолёты берут дальние участки, коптеры работают у своих баз в пределах "
              "радиуса действия."])

# 12. Сценарии
s = S[11]
ph = shapes_by_ph(s)
title(s, "ВСЕ ТИПЫ СЪЁМКИ")
shapes_by_ph(s)[0].width = Inches(8)
fill(ph[14], [
    ("RGB, мультиспектр, ИК — по GSD и перекрытиям; LiDAR — по плотности точек; геофизика — "
     "малая высота, огибание рельефа и секущие маршруты.", {}),
    ("Для LiDAR и геофизики в датасет добавлен Геоскан 401 (ТТХ из его руководства): в приложении "
     "к ТЗ таких нагрузок нет.", {}),
])
picture(s, 10, IMG / "geo.jpg")
picture(s, 11, IMG / "lidar.jpg")
picture(s, 12, IMG / "large.jpg")
notes(s, "Сверху справа — магнитная съёмка (секущие маршруты пунктиром), ниже — LiDAR, слева — 110 км².")

# 13. Масштабируемость
s = S[12]
ph = shapes_by_ph(s)
title(s, "МАСШТАБИРУЕМОСТЬ")
chart = next(sh for sh in s.shapes if sh.has_chart).chart
cd = CategoryChartData()
cd.categories = ["64 км²", "16 км²", "4 км²", "1 км²"]
cd.add_series("Время расчёта, с (8 бортов)", (4.94, 1.79, 1.78, 1.56))
chart.replace_data(cd)
chart.has_title = False
MID = dict(size=20, bold_first=True, bullets=False)
fill(ph[21], ["0,3–8 с"], **MID)
fill(ph[18], ["расчёт плана для областей до 64 км² и парка до 8 бортов"], bullets=False)
fill(ph[22], ["110 км²"], **MID)
fill(ph[23], ["план за 1,6 с, фронт Парето за 15 с"], bullets=False)
fill(ph[24], ["Дальность"], **MID)
fill(ph[25], ["крупные области делятся с учётом радиуса действия каждого борта"], bullets=False)
fill(ph[26], ["Docker"], **MID)
fill(ph[27], ["один контейнер, без внешних ключей; рельеф вшит при сборке"], bullets=False)
notes(s, "Диаграмма — время расчёта, секунды, парк из 8 бортов (201 и Gemini поровну).")

# 14. Уникальность
s = S[13]
ph = shapes_by_ph(s)
title(s, "УНИКАЛЬНОСТЬ")
fill(ph[14], bullets=False, paras=[
    ("Диспетчер парка, а не планировщик одного борта", {"bold": True, "size": 24}),
    ("Существующие программы строят задание одному аппарату. Сервис решает, кто что снимает, "
     "с какой базы и за сколько вылетов, и показывает цену выбора между скоростью и налётом.",
     {"size": 18, "space_before": 14}),
])
rows = [
    "Физика полёта: ветер, развороты самолёта, заряд, смена АКБ",
    "Самолёт и коптеры дополняют друг друга у запретных зон",
    "Разбиение с учётом радиуса действия каждого борта",
    "Проверка аварийного ухода на резервные площадки",
    "Каждая цифра ТТХ — со ссылкой на паспорт или пометкой «допущение»",
]
for i, r in enumerate(rows):
    fill(ph[15 + i], [r], bullets=False)

# 15. Архитектура и стек
s = S[14]
ph = shapes_by_ph(s)
title(s, "АРХИТЕКТУРА И СТЕК")
items = [
    ("Ядро", "Python 3.12: shapely, pyproj, networkx, numpy"),
    ("API", "FastAPI, OpenAPI-документация, pydantic"),
    ("Интерфейс", "React, TypeScript, Leaflet, OpenStreetMap"),
    ("Данные", "Датасет ТТХ в YAML, рельеф Copernicus DEM"),
    ("Экспорт", "GeoJSON (RFC 7946), KML 2.2, ZIP по плану"),
    ("Качество", "31 автотест, бенчмарк, Docker"),
]
for i, (h, d) in enumerate(items):
    fill(ph[49 + i], [f"0{i + 1}"])
    fill(ph[37 + 2 * i], [h], bullets=False)
    fill(ph[38 + 2 * i], [d], bullets=False)

# 16. Допущения и ограничения
s = S[15]
ph = shapes_by_ph(s)
title(s, "ДОПУЩЕНИЯ И ОГРАНИЧЕНИЯ")
lims = [
    "Часть ТТХ задана командой (АКБ 801, крейсерская скорость 201, смена АКБ) — помечены в датасете",
    "Энергия — по паспортной продолжительности с резервом и поправкой на ветер, без телеметрии",
    "Ветер постоянный по области и высоте; развороты самолёта без учёта сноса",
    "Одновременные полёты не разводятся по высоте; одна база на борт",
    "Оптимизация эвристическая: хороший план за секунды, глобальный оптимум не гарантирован",
]
for i, t in enumerate(lims):
    sh = ph[15 + i]
    sh.width = Inches(12.2)
    fill(sh, [t], bullets=False)

# 17. Развитие
s = S[16]
ph = shapes_by_ph(s)
title(s, "ПЛАНЫ РАЗВИТИЯ")
plans = [
    ("Точнее", "Калибровка энергомодели по телеметрии, прогноз ветра по высотам, трохоидальные развороты"),
    ("Безопаснее", "Эшелонирование бортов по высоте и времени, загрузка зон ограничений из официальных источников"),
    ("Ближе к полёту", "Экспорт в формат Geoscan Planner и MAVLink, перепланирование по ходу работ"),
]
for i, (h, d) in enumerate(plans):
    fill(ph[49 + i], [f"0{i + 1}"])
    fill(ph[37 + 2 * i], [h], bullets=False)
    fill(ph[38 + 2 * i], [d], bullets=False)

# 18. Финал
s = S[17]
ph = shapes_by_ph(s)
fill(ph[0], ["СПАСИБО!"])
fill(ph[1], bullets=False, paras=[
    ("Запуск", {"bold": True}),
    "docker compose up --build → http://127.0.0.1:8000",
    ("Код и документация", {"bold": True, "space_before": 10}),
    "[ссылка на репозиторий]",
    "README, docs/: архитектура, алгоритм, API, руководство, ограничения, бенчмарк",
    ("Контакты", {"bold": True, "space_before": 10}),
    "[капитан, Telegram]",
])

for n, sl in enumerate(S, 1):
    set_slide_number(sl, n)

prs.save(OUT)
print("saved", OUT, len(S), "slides")
